"""08_build_manuscript.py -- docx renderer over manuscript_content blocks.
Exposes render_docx(block_list, out_path) so the authored manuscript, the anonymized
manuscript and the title page can all be produced from the same renderer."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import manuscript_content as MC

FIG = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'results', 'figures')
OUTDIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'manuscript'); os.makedirs(OUTDIR, exist_ok=True)
NAVY = '1E2761'; SERIF = 'Times New Roman'

def render_docx(block_list, out_path):
    doc = Document()
    n = doc.styles['Normal']; n.font.name = SERIF; n.font.size = Pt(11)
    n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.15
    for h, sz in [('Heading 1', 14), ('Heading 2', 12)]:
        s = doc.styles[h]; s.font.name = SERIF; s.font.size = Pt(sz); s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(); fr.font.name = SERIF; fr.font.size = Pt(9)
    for t, v in [('begin', None), ('text', 'PAGE'), ('end', None)]:
        if t == 'text':
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = v
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t)
        fr._r.append(e)

    def para(text, align='justify', bold=False, italic=False, size=11, after=6, color=None):
        p = doc.add_paragraph()
        p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                       'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.name = SERIF; r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        return p

    def shade(cell, hexc):
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexc)
        cell._tc.get_or_add_tcPr().append(sh)

    def table(header, rows, widths, note=None, opts=None):
        opts = opts or {}
        t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ''; shade(c, NAVY)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h); r.bold = True; r.font.size = Pt(8); r.font.name = SERIF; r.font.color.rgb = RGBColor(255, 255, 255)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            boldrow = opts.get('bold_last_row') and ri == len(rows) - 1
            for i, v in enumerate(row):
                cells[i].text = ''; p = cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(str(v)); r.font.size = Pt(8); r.font.name = SERIF; r.bold = bool(boldrow)
        for r_ in t.rows:
            for i, c in enumerate(r_.cells):
                c.width = Inches(widths[i])
        if note:
            para(note, align='left', italic=True, size=8.5, after=10)

    def figure(fname, caption, width=6.3):
        doc.add_picture(os.path.join(FIG, fname), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(12)
        parts = caption.split('. ', 1)
        r1 = cp.add_run(parts[0] + '. '); r1.bold = True; r1.font.name = SERIF; r1.font.size = Pt(9)
        if len(parts) > 1:
            r2 = cp.add_run(parts[1]); r2.italic = True; r2.font.name = SERIF; r2.font.size = Pt(9)

    for blk in block_list:
        k = blk[0]
        if k == 'title':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(blk[1]); r.bold = True; r.font.name = SERIF; r.font.size = Pt(16)
            p.paragraph_format.space_after = Pt(10)
        elif k == 'authors':
            para(', '.join(blk[1]), align='center', size=11, after=2)
            para(blk[2], align='center', italic=True, size=9, after=12)
        elif k == 'subtitle':
            para(blk[1], align='center', italic=True, size=10, after=6)
        elif k == 'note':
            para(blk[1], align='center', italic=True, size=9.5, after=10)
        elif k == 'h_abstract':
            para(blk[1], bold=True, size=11, after=2)
        elif k == 'keywords':
            para(blk[1], italic=True, size=10, after=10)
        elif k == 'h1':
            doc.add_heading(blk[1], level=1)
        elif k == 'h2':
            doc.add_heading(blk[1], level=2)
        elif k == 'p':
            para(blk[1])
        elif k == 'table':
            header, rows, widths, note = blk[1], blk[2], blk[3], blk[4]
            opts = blk[5] if len(blk) > 5 else None
            table(header, rows, widths, note, opts)
        elif k == 'figure':
            figure(blk[1], blk[2])
        elif k == 'bullet':
            p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
            r = p.add_run(blk[1]); r.font.name = SERIF; r.font.size = Pt(10.5)
        elif k == 'refs':
            for r in blk[1]:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(r); run.font.name = SERIF; run.font.size = Pt(9.5)
    doc.save(out_path)
    print('Saved docx:', out_path, '| tables:', len(doc.tables), '| images:', len(doc.inline_shapes))

if __name__ == '__main__':
    render_docx(MC.blocks(), os.path.join(OUTDIR, 'AI_demand_planning_OOS_benchmark.docx'))
