"""build_highlights.py -- IJF/Elsevier Highlights file (3-5 bullets, <=85 chars each),
rendered to .docx, .md and .tex. Validates the length rule at build time."""
import os
from docx import Document
from docx.shared import Pt, Inches

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'manuscript')

HIGHLIGHTS = [
    'Three defensible evaluation choices each reverse the method ranking on the same data',
    'Croston rises from last to first at lead time; its sibling TSB stays mid-pack',
    'A same-week transacted price identifies sale weeks and flips the ML ranking',
    'Under RMSSE the naive benchmark goes from winning a fifth of series to worst',
    'Pretrained zero-shot models still lose to trained ones on every sparse class',
]
for h in HIGHLIGHTS:
    assert len(h) <= 85, f'{len(h)} chars: {h}'

# docx
doc = Document()
s = doc.styles['Normal']; s.font.name = 'Times New Roman'; s.font.size = Pt(11)
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
p = doc.add_paragraph(); r = p.add_run('Highlights'); r.bold = True; r.font.size = Pt(13)
for h in HIGHLIGHTS:
    bp = doc.add_paragraph(style='List Bullet'); br = bp.add_run(h)
    br.font.name = 'Times New Roman'; br.font.size = Pt(11)
doc.save(os.path.join(OUT, 'IJF_highlights.docx'))

# md
open(os.path.join(OUT, 'IJF_highlights.md'), 'w', encoding='utf-8').write(
    '# Highlights\n\n' + '\n'.join(f'- {h}' for h in HIGHLIGHTS) + '\n')

# tex
tex = ['\\documentclass[11pt]{article}', '\\usepackage[margin=1in]{geometry}',
       '\\usepackage[T1]{fontenc}\\usepackage{newtxtext}', '\\begin{document}',
       '\\section*{Highlights}', '\\begin{itemize}']
tex += [f'\\item {h}' for h in HIGHLIGHTS]
tex += ['\\end{itemize}', '\\end{document}']
open(os.path.join(OUT, 'IJF_highlights.tex'), 'w', encoding='utf-8').write('\n'.join(tex))
print('Saved IJF_highlights .docx/.md/.tex —', len(HIGHLIGHTS), 'bullets, all <=85 chars')
